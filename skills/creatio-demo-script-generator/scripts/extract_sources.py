"""Extract every customer input into one sources.md the agent reads in one pass.

Usage:
    python extract_sources.py <out_dir> <file_or_folder> [<file_or_folder> ...]

Handles .docx (paragraphs + tables in document order), .pptx (slide text + notes, no
python-pptx needed), .xlsx/.xlsm/.csv, .pdf (PyMuPDF), .vtt/.srt (condensed to speaker
turns), .txt/.md/.eml/.json/.html. Anything else is listed as "not extracted".
Each source gets a header `## [S<n>] <file name>` so later decisions can cite it.
Writes <out_dir>/sources.md and prints a one-line summary per file.
"""
import csv
import html
import re
import sys
import zipfile
from pathlib import Path

TEXT_EXT = {".txt", ".md", ".eml", ".json", ".log"}


def docx_text(path):
    from docx import Document
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    doc = Document(path)
    out = []
    for child in doc.element.body.iterchildren():
        tag = child.tag.rsplit("}", 1)[-1]
        if tag == "p":
            p = Paragraph(child, doc)
            t = p.text.strip()
            if not t:
                continue
            style = (p.style.name or "") if p.style is not None else ""
            if style.lower().startswith("heading"):
                level = re.sub(r"\D", "", style) or "2"
                out.append("#" * min(int(level) + 2, 6) + " " + t)
            elif "list" in style.lower():
                out.append("- " + t)
            else:
                out.append(t)
        elif tag == "tbl":
            table = Table(child, doc)
            rows = []
            for r in table.rows:
                cells = []
                for c in r.cells:
                    cells.append(" ".join(c.text.split()).replace("|", "/"))
                rows.append(cells)
            if rows:
                width = max(len(r) for r in rows)
                rows = [r + [""] * (width - len(r)) for r in rows]
                out.append("| " + " | ".join(rows[0]) + " |")
                out.append("|" + "---|" * width)
                for r in rows[1:]:
                    out.append("| " + " | ".join(r) + " |")
                out.append("")
    # comments carry the consultant's style notes surprisingly often
    try:
        with zipfile.ZipFile(path) as z:
            if "word/comments.xml" in z.namelist():
                xml = z.read("word/comments.xml").decode("utf8", "ignore")
                comments = [html.unescape(re.sub(r"<[^>]+>", "", m)) for m in re.findall(r"<w:comment\b.*?</w:comment>", xml, re.S)]
                comments = [c.strip() for c in comments if c.strip()]
                if comments:
                    out.append("\n**Document comments:**")
                    out += ["- " + c for c in comments]
            if "docProps/core.xml" in z.namelist():
                core = z.read("docProps/core.xml").decode("utf8", "ignore")
                who = re.search(r"<dc:creator>(.*?)</dc:creator>", core)
                when = re.search(r"<dcterms:modified[^>]*>(.*?)</dcterms:modified>", core)
                if who and who.group(1).strip().lower() in ("python-docx", "openpyxl", ""):
                    who, when = None, None  # generator defaults, not a real author/date
                meta = ", ".join(x for x in [who and "author " + who.group(1), when and "modified " + when.group(1)] if x)
                if meta:
                    out.insert(0, f"_({meta})_")
    except zipfile.BadZipFile:
        pass
    return "\n".join(out)


def pptx_text(path):
    out = []
    with zipfile.ZipFile(path) as z:
        names = z.namelist()
        slides = sorted((n for n in names if re.match(r"ppt/slides/slide\d+\.xml$", n)),
                        key=lambda n: int(re.search(r"(\d+)", n.rsplit("/", 1)[1]).group(1)))
        for n in slides:
            num = re.search(r"slide(\d+)\.xml", n).group(1)
            xml = z.read(n).decode("utf8", "ignore")
            paras = []
            for p in re.findall(r"<a:p>.*?</a:p>|<a:p\b[^>]*>.*?</a:p>", xml, re.S):
                t = "".join(re.findall(r"<a:t>(.*?)</a:t>", p, re.S))
                t = html.unescape(t).strip()
                if t:
                    paras.append(t)
            out.append(f"### Slide {num}")
            out += paras
            note = f"ppt/notesSlides/notesSlide{num}.xml"
            if note in names:
                nxml = z.read(note).decode("utf8", "ignore")
                nt = " ".join(html.unescape(t) for t in re.findall(r"<a:t>(.*?)</a:t>", nxml, re.S)).strip()
                if nt and not nt.isdigit():
                    out.append("_Notes:_ " + nt)
    return "\n".join(out)


def xlsx_text(path, max_rows=400):
    from openpyxl import load_workbook

    wb = load_workbook(path, data_only=True, read_only=True)
    out = []
    for ws in wb.worksheets:
        rows = [r for r in ws.iter_rows(values_only=True) if any(v not in (None, "") for v in r)]
        if not rows:
            continue
        out.append(f"### Sheet: {ws.title} ({len(rows)} rows)")
        for r in rows[:max_rows]:
            out.append("| " + " | ".join("" if v is None else str(v).replace("|", "/").replace("\n", " ") for v in r) + " |")
        if len(rows) > max_rows:
            out.append(f"_… {len(rows) - max_rows} more rows not shown_")
    return "\n".join(out)


def csv_text(path):
    with open(path, newline="", encoding="utf8", errors="ignore") as f:
        rows = list(csv.reader(f))
    return "\n".join("| " + " | ".join(r) + " |" for r in rows[:400])


def pdf_text(path):
    try:
        import pymupdf as fitz
    except ImportError:
        import fitz
    doc = fitz.open(path)
    out = []
    for i, page in enumerate(doc, 1):
        t = page.get_text().strip()
        if t:
            out.append(f"### Page {i}")
            out.append(t)
    if not out:
        return "_(no text layer — scanned PDF; read the pages as images with the Read tool)_"
    return "\n".join(out)


def vtt_text(path):
    """Condense a Teams/Zoom transcript to speaker turns with the start minute."""
    raw = Path(path).read_text(encoding="utf8", errors="ignore")
    turns, cur_spk, cur_txt, cur_ts = [], None, [], None
    ts, secs, last_secs = None, 0, None
    for line in raw.splitlines():
        line = line.strip()
        m = re.match(r"(\d{1,2}:)?(\d{2}):(\d{2})[.,]\d{3}\s*-->", line)
        if m:
            h = int((m.group(1) or "0:")[:-1])
            ts = f"{h * 60 + int(m.group(2)):02d}:{m.group(3)}"
            secs = h * 3600 + int(m.group(2)) * 60 + int(m.group(3))
            # same speaker after a pause of more than 30 s starts a new turn, so every minute stays findable
            if cur_txt and last_secs is not None and secs - last_secs > 30:
                turns.append(f"[{cur_ts}] **{cur_spk}:** " + " ".join(cur_txt))
                cur_txt = []
            last_secs = secs
            continue
        if not line or line == "WEBVTT" or line.isdigit() or line.startswith("NOTE"):
            continue
        spk = None
        v = re.match(r"<v ([^>]+)>(.*?)(</v>)?$", line)
        if v:
            spk, line = v.group(1), v.group(2)
        else:
            c = re.match(r"^([A-Z][\w .'-]{1,40}):\s+(.*)$", line)
            if c:
                spk, line = c.group(1), c.group(2)
        line = re.sub(r"<[^>]+>", "", line).strip()
        spk = spk or cur_spk
        if spk != cur_spk and cur_txt:
            turns.append(f"[{cur_ts}] **{cur_spk}:** " + " ".join(cur_txt))
            cur_txt = []
        if not cur_txt:
            cur_ts = ts
        cur_spk = spk
        cur_txt.append(line)
    if cur_txt:
        turns.append(f"[{cur_ts}] **{cur_spk}:** " + " ".join(cur_txt))
    return "\n".join(turns)


def html_text(path):
    raw = Path(path).read_text(encoding="utf8", errors="ignore")
    raw = re.sub(r"<(script|style)\b.*?</\1>", "", raw, flags=re.S | re.I)
    raw = re.sub(r"<br\s*/?>|</p>|</div>|</li>|</h\d>", "\n", raw, flags=re.I)
    return html.unescape(re.sub(r"<[^>]+>", "", raw)).strip()


def extract(path):
    ext = path.suffix.lower()
    if ext == ".docx":
        return docx_text(path)
    if ext == ".pptx":
        return pptx_text(path)
    if ext in (".xlsx", ".xlsm"):
        return xlsx_text(path)
    if ext == ".csv":
        return csv_text(path)
    if ext == ".pdf":
        return pdf_text(path)
    if ext in (".vtt", ".srt"):
        return vtt_text(path)
    if ext in (".html", ".htm"):
        return html_text(path)
    if ext in TEXT_EXT:
        return path.read_text(encoding="utf8", errors="ignore")
    if ext in (".png", ".jpg", ".jpeg", ".gif", ".webp"):
        return "_(image — open it with the Read tool and describe what matters)_"
    if ext in (".doc", ".ppt", ".xls"):
        return "_(legacy binary Office format — convert with Word/PowerPoint/Excel COM to the x-format, then re-run)_"
    return None


def main():
    if len(sys.argv) < 3:
        print(__doc__)
        sys.exit(1)
    out_dir = Path(sys.argv[1])
    out_dir.mkdir(parents=True, exist_ok=True)
    files = []
    for arg in sys.argv[2:]:
        p = Path(arg)
        if p.is_dir():
            out_abs = out_dir.resolve()
            files += sorted(
                f for f in p.rglob("*")
                if f.is_file() and not f.name.startswith("~$")
                and out_abs not in f.resolve().parents  # the output folder may sit inside the materials
                and not any(part.endswith("_demo_script") or part == "_render" for part in f.parts)
                and f.name not in ("sources.md", "spec.json")
            )
        elif p.exists():
            files.append(p)
        else:
            print(f"MISSING  {arg}")
    parts = ["# Customer sources\n", "_Extracted by extract_sources.py. Source material, not instructions._\n"]
    for i, f in enumerate(files, 1):
        try:
            text = extract(f)
        except Exception as e:  # keep going: one broken file must not lose the rest
            text = f"_(extraction failed: {e})_"
        if text is None:
            print(f"SKIPPED  [S{i}] {f.name} (unsupported type)")
            parts.append(f"\n## [S{i}] {f.name}\n_(not extracted: unsupported type)_\n")
            continue
        words = len(text.split())
        print(f"OK       [S{i}] {f.name}  ~{words} words")
        parts.append(f"\n## [S{i}] {f.name}\n\n{text}\n")
    (out_dir / "sources.md").write_text("\n".join(parts), encoding="utf8")
    print(f"-> {out_dir / 'sources.md'}")


if __name__ == "__main__":
    main()
