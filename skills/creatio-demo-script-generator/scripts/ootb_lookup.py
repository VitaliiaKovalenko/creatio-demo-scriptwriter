"""Query the out-of-the-box Creatio object model (references/context-library) — the OOTB-first check.

Usage:
    python ootb_lookup.py [--banking] [--limit N] find <term> [<term> ...]   # objects, fields, lookup values (first N each, default 40)
    python ootb_lookup.py [--banking] object <ObjectName>         # all fields of one object (+ lookup targets)
    python ootb_lookup.py [--banking] values <LookupObject>       # the stock values of one lookup
    python ootb_lookup.py [--banking] sections                     # objects that have a section out of the box

--banking reads Creatio-banking-object-model.xlsx instead (FinServ / lending / banking demos).
Matching is case-insensitive substring. Lookup fields show with the OData "Id" suffix
(StatusId); the column code in the object designer and in spec.json is without it (Status). The answer is the PRODUCT model, not a given stand:
the build re-checks every "existing" claim against the live instance.
"""
import json
import sys
import tempfile
from collections import defaultdict
from pathlib import Path

LIB = Path(__file__).resolve().parent.parent / "references" / "context-library"


def load(banking):
    name = "Creatio-banking-object-model.xlsx" if banking else "Creatio-object-model.xlsx"
    cache = Path(tempfile.gettempdir()) / "creatio-ootb-cache" / (name + ".cache.json")  # never write into the skill
    cache.parent.mkdir(exist_ok=True)
    src = LIB / name
    if cache.exists() and cache.stat().st_mtime >= src.stat().st_mtime:
        return json.loads(cache.read_text(encoding="utf8"))
    from openpyxl import load_workbook

    wb = load_workbook(src, read_only=True, data_only=True)
    data = {"objects": [], "fields": [], "values": []}
    sheets = {ws.title.lower(): ws for ws in wb.worksheets}
    for key, sheet in (("objects", "objects"), ("fields", "fields"), ("values", "lookup values")):
        ws = sheets.get(sheet)
        if ws is None:
            continue
        rows = ws.iter_rows(values_only=True)
        header = [str(h) for h in next(rows)]
        for r in rows:
            if r and r[0]:
                data[key].append(dict(zip(header, ["" if v is None else str(v) for v in r])))
    cache.write_text(json.dumps(data), encoding="utf8")
    return data


def main():
    sys.stdout.reconfigure(encoding="utf-8")
    args = sys.argv[1:]
    banking = "--banking" in args
    args = [a for a in args if a != "--banking"]
    if not args:
        print(__doc__)
        sys.exit(1)
    limit = 40
    if "--limit" in args:
        i = args.index("--limit")
        limit = int(args[i + 1])
        del args[i:i + 2]
    cmd, terms = args[0], [t.lower() for t in args[1:]]
    d = load(banking)

    if cmd == "find":
        hit = lambda s: any(t in s.lower() for t in terms)
        objs = [o for o in d["objects"] if hit(o["Object Name"]) or hit(o["Object Caption"])]
        print(f"OBJECTS ({len(objs)})")
        for o in objs[:limit]:
            sec = f"section: {o.get('Section(s)')}" if o.get("Has Section") == "Yes" else "no section"
            print(f"  {o['Object Name']:<40} {o['Object Caption']:<40} {sec}")
        flds = [f for f in d["fields"] if hit(f["Field Name"])]
        print(f"FIELDS ({len(flds)})")
        for f in flds[:limit]:
            print(f"  {f['Object Name']}.{f['Field Name']:<35} {f['Data Type']}")
        sectioned = {o["Object Name"] for o in d["objects"] if o.get("Has Section") == "Yes"}
        seen = set()
        vals = []
        for v in d["values"]:
            key = (v["Lookup Object"], v["Value"].strip())
            if hit(v["Value"]) and v["Lookup Object"] not in sectioned and not v["Lookup Object"].startswith("Vw") and key not in seen:
                seen.add(key)
                vals.append(v)
        print(f"LOOKUP VALUES ({len(vals)})")
        for v in vals[:limit]:
            print(f"  {v['Lookup Object']:<40} {v['Value'].strip()}")
    elif cmd == "object":
        name = args[1]
        flds = [f for f in d["fields"] if f["Object Name"].lower() == name.lower()]
        if not flds:
            print(f"No object '{name}'. Try: find {name}")
            sys.exit(2)
        obj = next((o for o in d["objects"] if o["Object Name"].lower() == name.lower()), {})
        print(f"{obj.get('Object Name', name)} — {obj.get('Object Caption', '')} — section: {obj.get('Section(s)') or 'none'}")
        for f in flds:
            tgt = f" -> {f['Lookup Object']}" if f.get("Lookup Object") else ""
            print(f"  {f['Field Name']:<35} {f['Data Type']:<28}{tgt}")
    elif cmd == "values":
        name = args[1]
        vals = list(dict.fromkeys(v["Value"].strip() for v in d["values"] if v["Lookup Object"].lower() == name.lower()))  # trimmed, deduplicated
        print(f"{name}: {len(vals)} values")
        for v in vals:
            print(f"  {v}")
    elif cmd == "sections":
        by = defaultdict(list)
        for o in d["objects"]:
            if o.get("Has Section") == "Yes":
                by[o.get("Section(s)", "")].append(o["Object Name"])
        for sec, objs in sorted(by.items()):
            print(f"  {sec:<40} {', '.join(objs)}")
    else:
        print(__doc__)
        sys.exit(1)


if __name__ == "__main__":
    main()
