"""spec.json -> the deliverables of a demo script package.

Usage:
    python build_outputs.py <spec.json> [--out <dir>] [--force] [--model-xlsx] [--docx]

Writes, next to the spec (or into --out) — documents as Markdown, Excel only for data:
    <Client>_Demo_Script.md       the script: 1 Brief customer summary, 2 Brief demo summary, 3 Industry
                                  context, 4 Roles and logins, 5 Demo use cases, 6 Competitive positioning
                                  (internal), Appendix (discovery insights, similar demos, supporting texts)
    <Client>_Demo_Script.html     the same script as one self-contained page, to publish as a shareable Artifact
                                  (build_html.py --external writes a copy without the internal sections)
    <Client>_Build_Spec.md        the SE's configuration: ground rules, build scope, before you present,
                                  AI actions, object model, lookups, pages, dashboards, DCM, processes,
                                  rules, integrations, notifications, clocks, demo data summary
    <Client>_Demo_Data.xlsx       one import-ready sheet per object + summary + lookup values
    <Client>_requirements.md      the build hand-off (creatio-demo-build phase 1 "derive" route)
    <Client>_Object_Model.xlsx    only with --model-xlsx: the same model as filterable sheets
    <Client>_*.docx               only with --docx: both documents in the Creatio Word template too
Runs validate_spec.validate() first and stops on errors unless --force.
"""
import json
import re
import sys
from collections import OrderedDict
from datetime import date, datetime
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))
from validate_spec import anchor_date, display_column, norm_type, render_date_tokens, resolve_date, validate  # noqa: E402
from quality_checks import tier_of, usage  # noqa: E402

# house style — references/docx-house-style.md
PRIMARY_800 = "123B54"
PRIMARY_600 = "1C6B8C"
TEAL = "0E8F86"
PRIMARY_300 = "8FC0D4"
PRIMARY_050 = "EFF6F9"
ROW_ALT = "F5F9FB"
HAIRLINE = "D8E2E8"
INK = "16212E"
SLATE = "51626E"
AMBER_BAR, AMBER_BG, AMBER_TXT = "C9861A", "FBF3E6", "8A530C"


def shorten(items, limit=220):
    """Join items with ', ' up to ~limit characters at an item boundary, then '+N more'."""
    out, n = [], 0
    for i, x in enumerate(items):
        if n + len(x) > limit and out:
            return ", ".join(out) + f" (+{len(items) - i} more)"
        out.append(x)
        n += len(x) + 2
    return ", ".join(out)


def apply_date_tokens(x, anchor):
    """Render {date:today-30} tokens in every string of the spec, so text dates move with meta.date."""
    if isinstance(x, dict):
        return {k: apply_date_tokens(v, anchor) for k, v in x.items()}
    if isinstance(x, list):
        return [apply_date_tokens(v, anchor) for v in x]
    if isinstance(x, str) and "{date:" in x:
        return render_date_tokens(x, anchor)
    return x


def safe_name(s):
    return re.sub(r"[^\w-]+", "_", s.strip()).strip("_")


# ======================================================================= DOCX
def build_docx(spec, path, part="script"):
    """Demo Script (part="script") or Build Specification (part="spec") in the Creatio template."""
    from doc_model import blocks
    from docx_render import render
    render(blocks(spec, part), path)


# ======================================================================= XLSX
def xlsx_styles():
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

    thin = Side(style="thin", color=HAIRLINE)
    return {
        "head_font": Font(name="Segoe UI", bold=True, color="FFFFFF", size=9),
        "head_fill": PatternFill("solid", fgColor=PRIMARY_600),
        "body_font": Font(name="Segoe UI", size=9, color=INK),
        "alt_fill": PatternFill("solid", fgColor=ROW_ALT),
        "main_fill": PatternFill("solid", fgColor=AMBER_BG),
        "border": Border(left=thin, right=thin, top=thin, bottom=thin),
        "wrap": Alignment(wrap_text=True, vertical="top"),
        "status": {
            "new": PatternFill("solid", fgColor="D5EFEC"),
            "extended": PatternFill("solid", fgColor="FBEBD0"),
            "existing": PatternFill("solid", fgColor="EDF0F2"),
        },
        "title_font": Font(name="Segoe UI Semibold", size=14, color=PRIMARY_800),
        "label_font": Font(name="Segoe UI", bold=True, size=9, color=PRIMARY_600),
    }


def write_sheet(wb, title, headers, rows, widths=None, status_col=None, main_rows=None, header_notes=None, number_formats=None):
    from openpyxl.comments import Comment
    from openpyxl.utils import get_column_letter

    S = xlsx_styles()
    ws = wb.create_sheet(title[:31])
    ws.append(headers)
    for i, h in enumerate(headers, 1):
        c = ws.cell(row=1, column=i)
        c.font, c.fill, c.border, c.alignment = S["head_font"], S["head_fill"], S["border"], S["wrap"]
        if header_notes and header_notes.get(h):
            c.comment = Comment(header_notes[h], "demo-script")
    for r_i, row in enumerate(rows, 2):
        ws.append(row)
        for c_i in range(1, len(headers) + 1):
            c = ws.cell(row=r_i, column=c_i)
            c.font, c.border, c.alignment = S["body_font"], S["border"], S["wrap"]
            if main_rows and (r_i - 2) in main_rows:
                c.fill = S["main_fill"]
            elif r_i % 2 == 1:
                c.fill = S["alt_fill"]
            if number_formats and headers[c_i - 1] in number_formats:
                c.number_format = number_formats[headers[c_i - 1]].get(r_i - 2, c.number_format) if isinstance(number_formats[headers[c_i - 1]], dict) else number_formats[headers[c_i - 1]]
        if status_col is not None:
            v = str(row[status_col]).lower()
            if v in S["status"]:
                ws.cell(row=r_i, column=status_col + 1).fill = S["status"][v]
    for i, h in enumerate(headers, 1):
        w = (widths or {}).get(h) if isinstance(widths, dict) else (widths[i - 1] if widths else None)
        if not w:
            longest = max([len(str(h))] + [len(str(r[i - 1])) for r in rows if i - 1 < len(r) and r[i - 1] is not None] or [10])
            w = min(max(10, longest + 2), 60)
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A2"
    if rows:
        ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{len(rows) + 1}"
    return ws


def summary_sheet(wb, title, heading, pairs):
    S = xlsx_styles()
    ws = wb.active
    ws.title = title
    ws["A1"] = heading
    ws["A1"].font = S["title_font"]
    r = 3
    for k, v in pairs:
        ws.cell(row=r, column=1, value=k).font = S["label_font"]
        c = ws.cell(row=r, column=2, value=v)
        c.font, c.alignment = S["body_font"], S["wrap"]
        r += 1
    ws.column_dimensions["A"].width = 26
    ws.column_dimensions["B"].width = 110
    return ws


def steps_str(x):
    return ", ".join(x.get("steps", []) or [])


def build_model_xlsx(spec, path):
    from openpyxl import Workbook

    meta, model = spec["meta"], spec.get("model", {})
    objs = model.get("objects", [])
    wb = Workbook()
    summary_sheet(wb, "Summary", f"{meta['client']} — Object model", [
        ("Version", f"{meta.get('version', 'v1')} · {meta.get('date', '')}"),
        ("Objects", f"{len(objs)} — {sum(o['status'] == 'new' for o in objs)} new, {sum(o['status'] == 'extended' for o in objs)} extended (new columns on an OOTB object), {sum(o['status'] == 'existing' for o in objs)} reused as is"),
        ("Columns", f"{sum(len(o.get('columns', [])) for o in objs)} listed — {sum(c.get('status') == 'new' for o in objs for c in o.get('columns', []))} new"),
        ("Lookups", f"{len(model.get('lookups', []))} — {sum(l['status'] == 'new' for l in model.get('lookups', []))} new"),
        ("Dashboards", str(len(model.get("dashboards", [])))),
        ("DCM cases", f"{len(model.get('dcm', []))} — a DCM whose stages carry steps is built by hand in the Case Designer"),
        ("Processes", f"{len(model.get('processes', []))} — {sum(p.get('build') == 'code' for p in model.get('processes', []))} buildable from code, {sum(p.get('build') == 'designer' for p in model.get('processes', []))} designer-only"),
        ("Business rules", str(len(model.get("rules", [])))),
        ("AI", str(len(model.get("ai", [])))),
        ("Status legend", "new = create in the working package · extended = OOTB object, new columns added via replacing object · existing = reuse as is"),
        ("Step references", "UC<n>.<step> — the step numbers of the Demo Script; they never renumber"),
        ("Caveat", "'existing' is checked against the product object model, not the stand. The build re-audits every claim against the live instance and its verdict wins."),
    ])
    write_sheet(wb, "Objects", ["Code", "Caption", "Status", "Section", "Purpose", "New columns", "Steps"],
                [[o["code"], o.get("caption", ""), o["status"], o.get("section", ""), o.get("purpose", ""),
                  sum(c.get("status") == "new" for c in o.get("columns", [])), steps_str(o)] for o in objs],
                widths=[22, 24, 10, 18, 50, 12, 30], status_col=2)
    write_sheet(wb, "Columns", ["Object", "Column code", "Caption", "Type", "Points to", "Status", "Required", "Steps", "Note"],
                [[o["code"], c["code"], c.get("caption", ""), c.get("type", ""), c.get("ref", "") or "", c.get("status", ""),
                  "yes" if c.get("required") else "", steps_str(c), c.get("note", "")] for o in objs for c in o.get("columns", [])],
                widths=[20, 24, 26, 14, 22, 10, 9, 26, 40], status_col=5)
    lk_rows = []
    for l in model.get("lookups", []):
        colors = l.get("colors", {})
        for i, v in enumerate(l.get("values", []), 1):
            lk_rows.append([l["code"], l.get("caption", ""), l["status"], i, v, colors.get(v, ""), steps_str(l) if i == 1 else ""])
    write_sheet(wb, "Lookups", ["Lookup", "Caption", "Status", "Order", "Value", "Color", "Steps"], lk_rows,
                widths=[24, 24, 10, 7, 32, 10, 30], status_col=2)
    write_sheet(wb, "Dashboards", ["Dashboard", "Location", "Widget", "Type", "Data source", "Example value", "Steps"],
                [[d.get("name", ""), d.get("location", ""), w.get("name", ""), w.get("type", ""), w.get("source", ""), w.get("example", ""), steps_str(d)]
                 for d in model.get("dashboards", []) for w in d.get("widgets", [])],
                widths=[24, 18, 26, 14, 50, 14, 24])
    dcm_rows = []
    for d in model.get("dcm", []):
        for i, s in enumerate(d.get("stages", []), 1):
            dcm_rows.append([d.get("name", d.get("object")), d.get("object"), d.get("stage_column"), i, s.get("name"),
                             "; ".join(s.get("elements", [])), "; ".join(s.get("exit_required", [])), steps_str(d) if i == 1 else ""])
    write_sheet(wb, "DCM", ["DCM case", "Object", "Stage column", "Order", "Stage", "Stage elements", "Required to leave", "Steps"], dcm_rows,
                widths=[22, 16, 14, 7, 22, 50, 30, 24])
    write_sheet(wb, "Processes", ["Process", "Trigger", "Actions (in order)", "Writes", "Build", "Steps", "Note"],
                [[p.get("name", ""), p.get("trigger", ""), "\n".join(f"{i}. {a}" for i, a in enumerate(p.get("actions", []), 1)),
                  p.get("writes", ""), p.get("build", ""), steps_str(p), p.get("note", "")] for p in model.get("processes", [])],
                widths=[28, 34, 60, 30, 10, 22, 30])
    write_sheet(wb, "Rules", ["Rule", "Object / page", "Level", "Trigger", "Condition", "Effect", "Steps"],
                [[r.get("name", ""), r.get("page") or r.get("object", ""), r.get("level", ""), r.get("trigger", ""), r.get("condition", ""), r.get("effect", ""), steps_str(r)]
                 for r in model.get("rules", [])], widths=[26, 24, 9, 22, 34, 34, 20])
    write_sheet(wb, "AI", ["Name", "Kind", "Purpose", "Inputs", "Output", "Steps"],
                [[a.get("name", ""), a.get("kind", ""), a.get("purpose", ""), a.get("inputs", ""), a.get("output", ""), steps_str(a)] for a in model.get("ai", [])],
                widths=[26, 20, 50, 30, 30, 20])
    write_sheet(wb, "Decisions", ["Kind", "Decision", "Why", "Where"],
                [[d.get("kind", "decision"), d.get("text", ""), d.get("why", ""), d.get("where", "")] for d in spec.get("decisions", [])],
                widths=[12, 60, 50, 16])
    write_sheet(wb, "Questions", ["#", "Where", "Question", "Default applied", "Who decides"],
                [[q["id"], q.get("where", ""), q["text"], q.get("default", ""), q.get("who", "")] for q in spec.get("questions", [])],
                widths=[6, 16, 60, 50, 14])
    wb.save(path)


def import_order(spec):
    """Objects ordered so that every lookup target among the seeded objects is loaded first."""
    model_objs = {o["code"]: o for o in spec.get("model", {}).get("objects", [])}
    blocks = OrderedDict((b["object"], b) for b in spec.get("data", []))
    deps = {}
    for oc in blocks:
        cols = model_objs.get(oc, {}).get("columns", [])
        deps[oc] = {c.get("ref") for c in cols if norm_type(c.get("type")) == "lookup" and c.get("ref") in blocks and c.get("ref") != oc}
    order, done = [], set()
    while len(order) < len(blocks):
        progress = False
        for oc in blocks:
            if oc not in done and deps[oc] <= done:
                order.append(oc)
                done.add(oc)
                progress = True
        if not progress:  # cycle: keep source order for the rest
            order += [oc for oc in blocks if oc not in done]
            break
    return order


def second_pass(spec, order):
    """Lookup columns whose target is imported later (a reference cycle): fill them in a second import pass."""
    pos = {oc: i for i, oc in enumerate(order)}
    model_objs = {o["code"]: o for o in spec.get("model", {}).get("objects", [])}
    late = []
    for oc in order:
        for c in model_objs.get(oc, {}).get("columns", []):
            ref = c.get("ref")
            if norm_type(c.get("type")) == "lookup" and ref in pos and pos[ref] > pos[oc]:
                late.append(f"{oc}.{c['code']}")
    return late


def build_data_xlsx(spec, path):
    from openpyxl import Workbook

    meta = spec["meta"]
    anchor = anchor_date(spec)
    model_objs = {o["code"]: o for o in spec.get("model", {}).get("objects", [])}
    lookups = spec.get("model", {}).get("lookups", [])
    blocks = {b["object"]: b for b in spec.get("data", [])}
    order = import_order(spec)
    late = second_pass(spec, order)
    # @key -> display value
    disp_of = {}
    for oc, b in blocks.items():
        dc = display_column(model_objs.get(oc, {}))
        for r in b.get("records", []):
            if r.get("_key"):
                disp_of[r["_key"]] = r.get(dc, r["_key"])

    total = sum(len(b.get("records", [])) for b in blocks.values())
    wb = Workbook()
    summary_sheet(wb, "Summary", f"{meta['client']} — Demo data", [
        ("Version", f"{meta.get('version', 'v1')} · {meta.get('date', '')}"),
        ("Records", f"{total} across {len(blocks)} objects"),
        ("Dates", f"Relative dates (today±N) resolved against {anchor.isoformat()}. Re-run build_outputs.py with a new meta.date to shift every date."),
        ("Main demo records", "Highlighted rows: the records the script names and shows on screen — spelled exactly as in the steps."),
        ("Import order", " → ".join(order) + "  (lookup targets first)"),
        ("Second pass", ", ".join(late) + " — reference cycle: fill these columns in a second import pass, after their targets exist" if late else "none"),
        ("Lookup values sheet", "Values to CREATE: every value of a new lookup, and only the new_values of an extended stock lookup."),
        ("Lookup columns", "Hold the display value of the target record or lookup value, as Creatio Excel import matches them."),
        ("Demo key", "Last column on each sheet: the record's id inside this package, used for cross-references. Import it into a "
                     "text column (or keep the sheet as the mapping): where a lookup cell holds a Demo key instead of a name — targets "
                     "whose names repeat, such as visits or order lines — match that column on Demo key, not by name. Header comments "
                     "on those columns say so."),
    ])
    write_sheet(wb, "Objects", ["Order", "Object", "Caption", "Sheet", "Records", "Main records", "Purpose"],
                [[i, oc, model_objs.get(oc, {}).get("caption", oc), (model_objs.get(oc, {}).get("caption") or oc)[:31],
                  len(blocks[oc].get("records", [])), sum(1 for r in blocks[oc].get("records", []) if r.get("_main")),
                  blocks[oc].get("purpose", "")] for i, oc in enumerate(order, 1)],
                widths=[7, 22, 24, 24, 9, 12, 60])
    lv_rows = [[l["code"], l.get("caption", ""), l["status"], v, l.get("colors", {}).get(v, "")]
               for l in lookups if l.get("status") in ("new", "extended")
               for v in (l.get("values", []) if l["status"] == "new" else l.get("new_values", []))]
    if lv_rows:
        write_sheet(wb, "Lookup values", ["Lookup", "Caption", "Status", "Value", "Color"], lv_rows,
                    widths=[24, 24, 10, 32, 10], status_col=2)
    used_titles = set()
    for oc in order:
        o = model_objs.get(oc, {})
        recs = blocks[oc].get("records", [])
        cols = [c for c in o.get("columns", []) if norm_type(c.get("type")) != "detail"]
        present = {k for r in recs for k in r if not k.startswith("_")}
        cols = [c for c in cols if c["code"] in present]
        headers = [c.get("caption") or c["code"] for c in cols] + ["Demo key"]
        notes = {(c.get("caption") or c["code"]): f"{oc}.{c['code']} · {c.get('type', '')}" + (f" → {c['ref']}" if c.get("ref") else "") for c in cols}
        rows, main, formats = [], set(), {}
        for i, r in enumerate(recs):
            row = []
            for c in cols:
                v = r.get(c["code"])
                t = norm_type(c.get("type"))
                if isinstance(v, str) and v.startswith("@") and model_objs.get(c.get("ref"), {}).get("allow_duplicate_names"):
                    notes[c.get("caption") or c["code"]] = notes.get(c.get("caption") or c["code"], "") + \
                        " · names repeat on the target: holds its Demo key — match on Demo key, not by name"
                    v = v[1:]
                if isinstance(v, str) and v.startswith("@"):
                    v = disp_of.get(v[1:], v)
                if t in ("date", "date/time"):
                    v = resolve_date(v, anchor)
                    if isinstance(v, str):
                        try:
                            v = datetime.fromisoformat(v) if t == "date/time" else date.fromisoformat(v[:10])
                        except ValueError:
                            pass
                    cap = c.get("caption") or c["code"]
                    formats[cap] = "yyyy-mm-dd hh:mm" if t == "date/time" else "yyyy-mm-dd"
                if t == "percentage":
                    formats[c.get("caption") or c["code"]] = "0.0"
                if t == "currency":
                    formats[c.get("caption") or c["code"]] = "#,##0.00"
                row.append(v)
            row.append(r.get("_key", ""))
            if r.get("_main"):
                main.add(i)
            rows.append(row)
        title = (o.get("caption") or oc)[:31]
        while title in used_titles or title in ("Summary", "Objects", "Lookup values"):
            title = (title[:28] + "_" + str(len(used_titles)))[:31]
        used_titles.add(title)
        write_sheet(wb, title, headers, rows, main_rows=main, header_notes=notes, number_formats=formats)
    wb.save(path)


# ======================================================================= REQUIREMENTS.MD
def build_requirements(spec, path, files):
    meta, model = spec["meta"], spec.get("model", {})
    L = []
    a = L.append
    a(f"# {meta['client']} — requirements (from the demo script package)\n")
    a(f"_Generated by creatio-demo-script-generator {meta.get('version', 'v1')} · {meta.get('date', '')}. "
      "Input for `creatio-demo-build` phase 1 — the 'derive' route is already done here; reconcile against the instance._\n")
    a("**Files:** " + " · ".join(f"`{f}`" for f in files) + "\n")
    a("## Scope\n")
    a("| UC | Name | Persona | Entry points | Steps |")
    a("|---|---|---|---|---|")
    for u in spec["use_cases"]:
        n = sum(1 for i in u.get("items", []) if "step" in i)
        a(f"| {u['id']} | {u['name']} | {u.get('persona', '')} | {', '.join(u.get('entry_points', []))} | {n} |")
    inst = meta.get("instance") or {}
    a(f"\n**Instance:** {inst.get('name', '')} {inst.get('url', '')} — mode **{inst.get('mode', '?')}**"
      + (f", reference demo {inst['reference_demo']}" if inst.get("reference_demo") else "")
      + (f". Previous demos: {inst['previous_demos']}" if inst.get("previous_demos") else "") + "\n")
    a("Credentials are in the team credential store — never in this file.\n")
    deal = meta.get("deal") or {}
    if isinstance(deal.get("arr_usd"), (int, float)):
        t = tier_of(deal["arr_usd"])
        a(f"**Deal tier:** {t['label']} — ARR {deal['arr_usd']:,.0f} USD" + (" (assumed)" if deal.get("arr_assumed") else "")
          + f". The tier sizes only the proposed additions ({t['proposals']['min']}–{t['proposals']['max']}, up to effort "
          f"{t['proposals']['max_effort']}); the agreed configuration is not capped.\n")
    if spec.get("roles"):
        a("## Roles and logins\n")
        a("| Role | Login mode | Login | Workplace | Use cases |")
        a("|---|---|---|---|---|")
        for r in spec["roles"]:
            a(f"| {r.get('name', '')} | {r.get('login_mode', '')} | {r.get('login', '')} | {r.get('workplace', '')} | {', '.join(r.get('use_cases', []))} |")
    if spec.get("ground_rules"):
        a("\n## Ground rules\n")
        for g in spec["ground_rules"]:
            a(f"- {g}")
    prio = sorted([u for u in spec["use_cases"] if u.get("priority")], key=lambda u: {"must": 0, "should": 1, "could": 2}.get(u["priority"], 3))
    if prio:
        a("\n**Protect in order if time runs short:** " + " → ".join(f"{u['id']} ({u['priority']})" for u in prio)
          + (f" · budget {meta['build_budget_h']} h" if meta.get("build_budget_h") else "") + "\n")
    if spec.get("out_of_scope"):
        a("## Out of scope — do not build\n")
        for x in spec["out_of_scope"]:
            a(f"- {x.get('item', '')} — {x.get('why', '')}")
    pre = spec.get("pre_demo", {})
    if pre:
        a("\n## Data state before the demo\n")
        for x in pre.get("must_exist", []):
            a(f"- exists: {x}")
        for x in pre.get("must_not_exist", []):
            a(f"- must NOT exist (created live): {x}")
        for x in pre.get("checks", []):
            a(f"- check: {x}")
    a("")
    a("## Objects\n")
    a("| Object | Caption | Action | Steps |")
    a("|---|---|---|---|")
    act = {"new": "NEW custom object (+ section, list and form pages)", "extended": "extend base object (replacing)", "existing": "reuse as is"}
    for o in model.get("objects", []):
        a(f"| `{o['code']}` | {o.get('caption', '')} | {act.get(o['status'], o['status'])} | {steps_str(o)} |")
    for o in model.get("objects", []):
        new_cols = [c for c in o.get("columns", []) if c.get("status") == "new"]
        if not new_cols:
            continue
        a(f"\n### {o['code']} — new columns\n")
        a("| Column | Caption | Type | Points to | Required | Steps |")
        a("|---|---|---|---|---|---|")
        for c in new_cols:
            a(f"| `{c['code']}` | {c.get('caption', '')} | {c.get('type', '')} | {c.get('ref') or ''} | {'yes' if c.get('required') else ''} | {steps_str(c)} |")
    a("\n## Lookups\n")
    users = {}
    for o in model.get("objects", []):
        for c in o.get("columns", []):
            if c.get("ref"):
                users.setdefault(c["ref"], []).append(o["code"])
    for l in model.get("lookups", []):
        shared = users.get(l["code"], [])
        tag = f" — shared by {', '.join(sorted(set(shared)))}" if len(set(shared)) > 1 else ""
        a(f"- `{l['code']}` ({l['status']}){tag}: " + " / ".join(f'"{v}"' for v in l.get("values", [])))
    if model.get("dcm"):
        a("\n## DCM — designer-only when stages carry steps\n")
        for d in model["dcm"]:
            a(f"- **{d.get('name', d['object'])}** on `{d['object']}`.`{d.get('stage_column')}`: " + " → ".join(s["name"] for s in d.get("stages", [])) + f"  ({steps_str(d)})")
    if model.get("processes"):
        a("\n## Processes\n")
        a("| Process | Trigger | Does | Real? | Audience sees | Build | Steps |")
        a("|---|---|---|---|---|---|---|")
        for p in model["processes"]:
            sim = f" (simulated: {p['simulation']})" if p.get("simulation") else ""
            a(f"| {p.get('name', '')} | {p.get('trigger', '')} | {' → '.join(p.get('actions', []))} | {p.get('realization', '')}{sim} | "
              f"{p.get('audience_sees', '')} | {'buildable from code' if p.get('build') == 'code' else 'designer-only'} | {steps_str(p)} |")
    if model.get("rules"):
        a("\n## Business rules\n")
        for r in model["rules"]:
            a(f"- **{r.get('name', '')}** ({r.get('level', '')}, {r.get('page') or r.get('object', '')}): when {r.get('condition', '')} → {r.get('effect', '')}  ({steps_str(r)})")
    if model.get("ui"):
        a("\n## Pages, buttons and navigation\n")
        a("| Element | Kind | Where | Does | Steps |")
        a("|---|---|---|---|---|")
        for x in model["ui"]:
            a(f"| {x.get('name', '')} | {x.get('kind', '')} | {x.get('where', '')} | {x.get('does', '')} | {steps_str(x)} |")
    if model.get("dashboards"):
        a("\n## Dashboards\n")
        for d in model["dashboards"]:
            a(f"- **{d.get('name', '')}** ({d.get('location', '')}): " + "; ".join(f"{w.get('name')} [{w.get('type')}]" for w in d.get("widgets", [])))
    if model.get("ai"):
        a("\n## AI actions\n")
        for x in model["ai"]:
            a(f"- **{x.get('name', '')}** ({x.get('kind', '')}, {x.get('realization', '')}) — {x.get('purpose', '')}  ({steps_str(x)})")
            a(f"  - trigger: {x.get('trigger', '')} · input: {x.get('inputs', '')} · output: {x.get('output', '')}")
            a(f"  - stored in: {x.get('stored_in', '')} · review: {x.get('review', '')} · audience sees: {x.get('audience_sees', '')}")
            if x.get("example_output"):
                a(f"  - example output: \"{x['example_output']}\"")
    if model.get("integrations"):
        a("\n## Integrations\n")
        for i in model["integrations"]:
            a(f"- **{i.get('name', '')}** stands for {i.get('system', '')} — {i.get('realization', '')}: {i.get('pattern') or i.get('simulation', '')}; "
              f"populates {i.get('populates', '')}" + (f"; timeline note \"{i['timeline_note']}\"" if i.get("timeline_note") else "") + f"  ({steps_str(i)})")
    if model.get("notifications"):
        a("\n## Notifications\n")
        for n in model["notifications"]:
            a(f"- **{n.get('id', '')} {n.get('name', '')}** — {n.get('trigger', '')} → {n.get('recipient', '')} via {n.get('channel', '')} ({n.get('realization', '')}): {n.get('copy', '')}")
    if model.get("slas"):
        a("\n## SLAs and clocks\n")
        for s_ in model["slas"]:
            a(f"- **{s_.get('name', '')}**: {s_.get('duration', '')}, starts {s_.get('start', '')}, stops {s_.get('stop', '')}, shown in {s_.get('shown_in', '')}")
    a("\n## Demo data\n")
    for oc in import_order(spec):
        b = next(b for b in spec["data"] if b["object"] == oc)
        main = sum(1 for r in b.get("records", []) if r.get("_main"))
        a(f"- `{oc}`: {len(b.get('records', []))} records ({main} main) — {b.get('purpose', '')}")
    a("\n## Decisions taken\n")
    for d in spec.get("decisions", []):
        a(f"- [{d.get('kind', 'decision')}] {d.get('text', '')} — {d.get('why', '')} ({d.get('where', '')})")
    a("\n## Open questions (default applied)\n")
    for q in spec.get("questions", []):
        a(f"- **{q['id']}** ({q.get('where', '')}, {q.get('who', '')}): {q['text']} → default: {q.get('default', '')}")
    a("\n## IMPLEMENTATION STATUS\n\n_(the build appends here)_\n")
    Path(path).write_text("\n".join(L), encoding="utf8")


def main():
    sys.stdout.reconfigure(encoding="utf-8")
    args = sys.argv[1:]
    force = "--force" in args
    model_xlsx = "--model-xlsx" in args
    want_docx = "--docx" in args
    out = None
    if "--out" in args:
        i = args.index("--out")
        out = Path(args[i + 1])
        del args[i:i + 2]
    args = [x for x in args if not x.startswith("--")]
    if len(args) != 1:
        print(__doc__)
        sys.exit(2)
    spec_path = Path(args[0])
    out = out or spec_path.parent
    out.mkdir(parents=True, exist_ok=True)
    spec = json.loads(spec_path.read_text(encoding="utf8"))
    errors, warnings = validate(spec)
    for e in errors:
        print("ERROR  ", e)
    print(f"{len(errors)} errors, {len(warnings)} warnings (run validate_spec.py to see the warnings)")
    if errors and not force:
        print("Not building — fix the errors (or --force for a draft).")
        sys.exit(1)
    base = safe_name(spec["meta"]["client"])
    data_x, req = f"{base}_Demo_Data.xlsx", f"{base}_requirements.md"
    spec = apply_date_tokens(spec, anchor_date(spec))
    from build_markdown import build_markdown  # the documents are Markdown: read in any editor, on GitHub, diff cleanly
    script_md, build_spec_md = f"{base}_Demo_Script.md", f"{base}_Build_Spec.md"
    build_markdown(spec, out / script_md, part="script")
    build_markdown(spec, out / build_spec_md, part="spec")
    from build_html import build_html  # the shareable page: publish it as an Artifact, same blocks as the .md
    script_html = f"{base}_Demo_Script.html"
    build_html(spec, out / script_html)
    build_data_xlsx(spec, out / data_x)
    files = [script_md, script_html, build_spec_md, data_x]
    if want_docx:  # Word only on request; the same blocks, so the formats never drift apart
        script, build_spec = f"{base}_Demo_Script.docx", f"{base}_Build_Spec.docx"
        build_docx(spec, out / script, part="script")
        build_docx(spec, out / build_spec, part="spec")
        files += [script, build_spec]
    if model_xlsx:
        files.append(f"{base}_Object_Model.xlsx")
        build_model_xlsx(spec, out / files[-1])
    build_requirements(spec, out / req, files + [spec_path.name])
    for f in files + [req]:
        print("->", out / f)


if __name__ == "__main__":
    main()
